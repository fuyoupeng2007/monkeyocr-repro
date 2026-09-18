#!/usr/bin/env python3
"""Render each slide of the generated PPTX into a PNG for visual verification.

There is no LibreOffice on the box, so instead of rasterising the pptx directly
we re-draw every slide's shapes with Pillow using a CJK font (wqy-zenhei), which
verifies: (a) the slide geometry/overflow, (b) that CJK text is present and
renderable, (c) table dimensions. Outputs go to outputs/deliverables/_slides/.
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

DEL = Path("/root/autodl-tmp/monkeyocr-repro/outputs/deliverables")
OUTDIR = DEL / "_slides"
OUTDIR.mkdir(parents=True, exist_ok=True)
CJK = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"

EMU_PER_PX = 9525  # 96 dpi


def font(size_pt: float):
    px = max(10, int(size_pt * 96 / 72))
    try:
        return ImageFont.truetype(CJK, px)
    except Exception:
        return ImageFont.load_default()


def main() -> None:
    prs = Presentation(str(DEL / "汇报材料.pptx"))
    W = int(prs.slide_width / EMU_PER_PX)
    H = int(prs.slide_height / EMU_PER_PX)
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            x, y, w, h = (int(v / EMU_PER_PX) for v in (sh.left, sh.top, sh.width, sh.height))
            if sh.has_text_frame and sh.text_frame.text.strip():
                yy = y
                for p in sh.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    if not txt.strip():
                        continue
                    size = 18
                    for r in p.runs:
                        if r.font.size:
                            size = r.font.size.pt
                            break
                    f = font(size)
                    d.multiline_text((x, yy), txt, fill="black", font=f)
                    bbox = d.multiline_textbbox((x, yy), txt, font=f)
                    yy = bbox[3] + 6
            if getattr(sh, "has_table", False) and sh.has_table:
                tbl = sh.table
                nrow, ncol = len(tbl.rows), len(tbl.columns)
                cw, ch = w // max(1, ncol), h // max(1, nrow)
                for i in range(nrow):
                    for j in range(ncol):
                        cx, cy = x + j * cw, y + i * ch
                        d.rectangle([cx, cy, cx + cw, cy + ch], outline="#bbbbbb")
                        d.multiline_text((cx + 4, cy + 4), tbl.cell(i, j).text[:22], fill="black", font=font(11))
        img.save(OUTDIR / f"slide_{idx:02d}.png")
        print(f"rendered slide_{idx:02d}.png {W}x{H}")


if __name__ == "__main__":
    main()
