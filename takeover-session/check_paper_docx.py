import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

path = r"D:\haeness\monkeyocr-takeover\deliverables\MonkeyOCR论文_中文全译.docx"
d = Document(path)

heads = [(p.style.name, p.text) for p in d.paragraphs if p.style.name.startswith("Heading")]
print("标题层级：")
for s, t in heads:
    print(f"  {s:9s} {t[:60]}")

print("\n表格：")
for i, t in enumerate(d.tables, 1):
    print(f"  表{i}: {len(t.rows)} 行 × {len(t.columns)} 列 | 首行: "
          + " | ".join(c.text[:12] for c in t.rows[0].cells)[:90])

chars = sum(len(p.text) for p in d.paragraphs)
tchars = sum(len(c.text) for t in d.tables for r in t.rows for c in r.cells)
print(f"\n正文字符: {chars}  表格字符: {tchars}  合计: {chars + tchars}")

# font check
for p in d.paragraphs:
    if p.runs:
        from docx.oxml.ns import qn
        rPr = p.runs[0]._element.find(qn("w:rPr"))
        if rPr is not None:
            rf = rPr.find(qn("w:rFonts"))
            if rf is not None:
                print("东亚字体:", rf.get(qn("w:eastAsia")), "| 西文字体:", rf.get(qn("w:ascii")))
                break
