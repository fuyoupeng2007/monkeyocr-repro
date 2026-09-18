#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the email-to-professor markdown into a clean Word document.

Layout choices:
  * the email body is the main content, in a readable 11pt with 1.5 line spacing
    (so it can be copied straight into a mail client without reformatting);
  * a distinct '抬头区块' at the top carries 收件人/主题/附件;
  * everything after the separator (发送前核对清单, 设计说明) is an appendix in a
    slightly muted style, since it is for the sender, not the recipient.

All string literals use single quotes so Chinese prose can use double quotes.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / 'deliverables' / '给刘禹良老师的邮件.md'
DST = HERE / 'deliverables' / '给刘禹良老师的邮件.docx'

CJK = '微软雅黑'
ACCENT = RGBColor(0x16, 0x68, 0xDC)
MUTED = RGBColor(0x5B, 0x64, 0x72)
WARN = RGBColor(0xA8, 0x07, 0x1A)


def set_run(run, size=10.5, bold=False, color=None, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = CJK
    if color is not None:
        run.font.color.rgb = color
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:ascii'), CJK)
    rf.set(qn('w:hAnsi'), CJK)
    rf.set(qn('w:eastAsia'), CJK)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def add_inline(p, text, size=10.5, bold=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            set_run(p.add_run(part[2:-2]), size=size, bold=True, color=color)
        elif part.startswith('`') and part.endswith('`'):
            set_run(p.add_run(part[1:-1]), size=size - 0.5, bold=bold, color=color)
        else:
            set_run(p.add_run(part), size=size, bold=bold, color=color)


def para(doc, text='', size=11, bold=False, color=None, align=None,
         space_before=0, space_after=8, indent=0, line=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if text:
        add_inline(p, text, size=size, bold=bold, color=color)
    return p


def bullet(doc, text, size=10.5, color=None):
    p = doc.add_paragraph(style='List Bullet')
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.45
    add_inline(p, text, size=size, color=color)
    return p


def h(doc, text, level=1):
    size = {1: 15, 2: 12.5}.get(level, 11.5)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(6)
    set_run(p.add_run(text), size=size, bold=True, color=ACCENT if level == 1 else None)
    if level == 1:
        pPr = p._element.get_or_add_pPr()
        bdr = pPr.makeelement(qn('w:pBdr'), {})
        bdr.append(bdr.makeelement(qn('w:bottom'), {
            qn('w:val'): 'single', qn('w:sz'): '8', qn('w:space'): '1',
            qn('w:color'): '1668DC'}))
        pPr.append(bdr)
    return p


def info_block(doc, lines):
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in lines:
        cells = t.add_row().cells
        cells[0].text = ''
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        set_run(p0.add_run(label), size=10, bold=True)
        cells[1].text = ''
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        add_inline(p1, value, size=10)
    for row in t.rows:
        row.cells[0].width = Cm(2.6)
        row.cells[1].width = Cm(13.4)
    return t


def table(doc, rows, font=9.5):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ''
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, str(val), size=font, bold=(i == 0))
    return t


def separator(doc):
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    bdr = pPr.makeelement(qn('w:pBdr'), {})
    bdr.append(bdr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single', qn('w:sz'): '6', qn('w:space'): '1', qn('w:color'): 'C0C4CC'}))
    pPr.append(bdr)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def main() -> None:
    lines = SRC.read_text(encoding='utf-8').splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    st = doc.styles['Normal']
    st.font.name = CJK
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)

    # ---- title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run('致刘禹良老师的科研申请邮件'), size=18, bold=True)
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(12)
    set_run(q.add_run('人工智能 2505 班　傅有鹏　｜　2026-09-18'), size=10, color=MUTED)

    # ---- parse: header table (收件人/主题/附件) then email body, then appendix
    i = 0
    header: list[tuple[str, str]] = []
    body_started = False
    appendix_started = False
    in_table = False
    tbl: list[list[str]] = []

    def flush():
        nonlocal tbl, in_table
        if tbl:
            table(doc, tbl)
            tbl = []
        in_table = False

    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()

        # drop the file's own H1 — the Word title already says the same thing
        if s.startswith('# ') and not body_started:
            i += 1
            continue

        # collect the info block at the top
        m = re.match(r'^\*\*(收件人|主题|附件建议)\*\*[：:]\s*(.*)$', s)
        if m and not body_started:
            header.append((m.group(1), m.group(2)))
            i += 1
            continue

        if not body_started and s == '---':
            body_started = True
            info_block(doc, header)
            para(doc, '', space_after=2)
            i += 1
            continue

        if body_started and not appendix_started and s == '---':
            appendix_started = True
            flush()
            doc.add_page_break()
            i += 1
            continue

        # markdown table
        if s.startswith('|') and s.endswith('|'):
            cells = [c.strip() for c in s.strip('|').split('|')]
            if set(''.join(cells)) <= set('-: '):
                i += 1
                continue
            tbl.append(cells)
            in_table = True
            i += 1
            continue
        if in_table:
            flush()

        if not s:
            i += 1
            continue

        if s.startswith('## '):
            h(doc, s[3:].strip(), 1 if appendix_started else 2)
            i += 1
            continue
        if s.startswith('# '):
            h(doc, s[2:].strip(), 1)
            i += 1
            continue
        if s.startswith('- '):
            bullet(doc, s[2:].strip(), size=10.5 if appendix_started else 11,
                   color=MUTED if appendix_started else None)
            i += 1
            continue

        # body paragraph: keep quotes as-is, 11pt for the email itself
        if body_started and not appendix_started:
            para(doc, s, size=11, space_after=8)
        else:
            para(doc, s, size=10, color=MUTED if appendix_started else None, space_after=6)
        i += 1

    if in_table:
        flush()

    doc.save(str(DST))
    print('written:', DST.name, DST.stat().st_size, 'bytes')
    d = Document(str(DST))
    print('段落:', len(d.paragraphs), '表格:', len(d.tables))


if __name__ == '__main__':
    main()
