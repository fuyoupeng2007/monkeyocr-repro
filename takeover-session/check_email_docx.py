import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
p = r"D:\haeness\monkeyocr-takeover\deliverables\给刘禹良老师的邮件.docx"
d = Document(p)

print("=== 抬头信息表 ===")
for t in d.tables:
    for row in t.rows:
        cells = [c.text.strip() for c in row.cells]
        if any(cells):
            print("  ", " | ".join(cells))

print("\n=== 正文段落（前 12 段）===")
n = 0
for para in d.paragraphs:
    if para.text.strip():
        n += 1
        if n <= 12:
            print(f"  [{n}] {para.text[:95]}")
print(f"  ... 共 {n} 段非空")

print("\n=== 关键内容检查 ===")
full = "\n".join(x.text for x in d.paragraphs) + "\n".join(
    c.text for t in d.tables for r in t.rows for c in r.cells)
for k in ["刘老师", "ylliu@hust.edu.cn", "管海粟", "ACL 2024", "最佳论文",
          "MonkeyOCR", "SRR", "0.2022", "0.8081", "0.1442", "揭榜挂帅", "伍冬睿",
          "AI 编程工具", "86.41", "62/170", "8 到 12 小时"]:
    c = full.count(k)
    flag = "OK " if c else "!! "
    print(f"  {flag}{k}: {c}")

print("\n=== 不应出现的内容（我上次写错的）===")
for k in ["伍老师", "伍东睿", "伍冬睿老师课题组下"]:
    print(f"  {k}: {full.count(k)}")

print("\n字符总数:", len(full))
