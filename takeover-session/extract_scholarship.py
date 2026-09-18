"""Extract the scholarship materials (PPT / speech / application form) to text."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from pptx import Presentation

BASE = Path(r"D:\haeness\奖学金")
OUT = Path(r"D:\haeness\monkeyocr-takeover\scholarship_text.txt")
sys.stdout.reconfigure(encoding="utf-8")

chunks: list[str] = []

ppt = BASE / "三奖答辩_傅有鹏.pptx"
if ppt.exists():
    prs = Presentation(str(ppt))
    chunks.append("=" * 30 + " PPT 三奖答辩 " + "=" * 30)
    for i, slide in enumerate(prs.slides, 1):
        chunks.append(f"\n--- 幻灯片 {i} ---")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                chunks.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    chunks.append(" | ".join(c.text.strip() for c in row.cells))
    print(f"PPT: {len(prs.slides)} slides")

for name in ["三奖答辩_演讲稿_傅有鹏.docx", "奖学金申报表—傅有鹏(1).docx",
             "附件1：2025级奖学金申请审批表—傅有鹏.docx"]:
    f = BASE / name
    if not f.exists():
        continue
    d = Document(str(f))
    chunks.append("\n" + "=" * 30 + f" {name} " + "=" * 30)
    for p in d.paragraphs:
        if p.text.strip():
            chunks.append(p.text.strip())
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))
    print(f"{name}: {len(d.paragraphs)} paragraphs, {len(d.tables)} tables")

text = "\n".join(chunks)
OUT.write_text(text, encoding="utf-8")
print("written:", OUT, len(text), "chars")
