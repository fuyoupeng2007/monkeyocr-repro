import sys
from docx import Document
from docx.oxml.ns import qn

for path in sys.argv[1:]:
    d = Document(path)
    paras = d.paragraphs
    tables = d.tables
    chars = sum(len(p.text) for p in paras) + sum(
        len(c.text) for t in tables for r in t.rows for c in r.cells)
    heads = [p.text for p in paras if p.runs and p.runs[0].font.size and p.runs[0].font.size.pt >= 13]
    print("=" * 70)
    print("file   :", path.split("\\")[-1])
    print("段落   :", len(paras), " 表格:", len(tables), " 字符数:", chars)
    print("章节   :", " / ".join(heads))
    # eastAsia font check on the first body run
    checked = 0
    for p in paras:
        for r in p.runs:
            rPr = r._element.find(qn("w:rPr"))
            if rPr is not None:
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None:
                    print("东亚字体:", rf.get(qn("w:eastAsia")), " 西文字体:", rf.get(qn("w:ascii")))
                    checked = 1
                    break
        if checked:
            break
    # sample content
    samples = [p.text for p in paras if p.text.strip()][:3]
    for s in samples:
        print("   样本:", s[:60])
