import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
p = r"D:\haeness\monkeyocr-takeover\deliverables\给导师的汇报_MonkeyOCR论文复现_v2.docx"
d = Document(p)

print("=== 章节 ===")
for para in d.paragraphs:
    if para.runs and para.runs[0].font.size and para.runs[0].font.size.pt >= 13:
        print("  ", para.text)

print("\n=== 表格 ===")
for i, t in enumerate(d.tables, 1):
    print(f"  表{i}: {len(t.rows)}行x{len(t.columns)}列 | 首行: "
          + " / ".join(c.text[:14] for c in t.rows[0].cells)[:95])

print("\n=== AI 归因检查 ===")
full = "\n".join(x.text for x in d.paragraphs) + "\n".join(
    c.text for t in d.tables for r in t.rows for c in r.cells)
for k in ["Codex", "DSH", "AI 编程代理", "本人主导", "本人验收", "工程支撑"]:
    print(f"  '{k}': {full.count(k)} 次")

print("\n=== 关键数字核对 ===")
for k in ["0.2022", "0.8081", "0.2363", "0.4136", "0.1442", "0.058", "0.5741", "0.6064",
          "81", "39", "118", "120"]:
    print(f"  '{k}': {full.count(k)} 次")

print("\n=== 字体 ===")
for para in d.paragraphs:
    if para.runs:
        from docx.oxml.ns import qn
        rPr = para.runs[0]._element.find(qn("w:rPr"))
        if rPr is not None and rPr.find(qn("w:rFonts")) is not None:
            rf = rPr.find(qn("w:rFonts"))
            print("  东亚字体:", rf.get(qn("w:eastAsia")), "| 西文:", rf.get(qn("w:ascii")))
            break

print("\n字符总数:", len(full))
