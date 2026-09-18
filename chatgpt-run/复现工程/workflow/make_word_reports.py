from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

ROOT = Path(r"D:\chatgpt\刘禹良团队_入组调研与复现审批\复现工程\deliverables")

def add_table(doc, lines):
    rows = [[c.strip() for c in x.strip().strip('|').split('|')] for x in lines if not set(x.replace('|','').strip()) <= {'-'}]
    if not rows: return
    table = doc.add_table(rows=len(rows), cols=len(rows[0])); table.style='Table Grid'
    for i,row in enumerate(rows):
        for j,cell in enumerate(row):
            cell = table.cell(i,j); cell.text = row[j] if j < len(row) else ''
            for p in cell.paragraphs:
                for r in p.runs: r.font.size=Pt(9.5)
                if i==0:
                    for r in p.runs: r.font.bold=True

def convert(src, dst):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name='微软雅黑';normal.font.size=Pt(10.5)
    lines=src.read_text(encoding='utf-8').splitlines(); i=0
    while i<len(lines):
        s=lines[i]
        if s.startswith('# '):
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(s[2:]).font.color.rgb=RGBColor(0,0,0)
        elif s.startswith('## '): doc.add_heading(s[3:], level=1)
        elif s.startswith('### '): doc.add_heading(s[4:], level=2)
        elif s.startswith('|'):
            buf=[]
            while i<len(lines) and lines[i].startswith('|'): buf.append(lines[i]);i+=1
            add_table(doc,buf);continue
        elif s.strip():
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
            if s.startswith('**') and s.endswith('**'): p.add_run(s.strip('*')).bold=True
            else: p.add_run(s)
        i+=1
    doc.save(dst)

convert(ROOT/'导师提交附件_ MonkeyOCR复现阶段汇报.md', ROOT/'导师提交附件 MonkeyOCR复现阶段汇报.docx')
convert(ROOT/'自己讲懂的版本_ MonkeyOCR复现.md', ROOT/'自己讲懂的版本 MonkeyOCR复现.docx')
